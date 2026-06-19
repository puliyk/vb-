#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把生成的 PNG 嵌入 Word，替换占位文字"""
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

D = r'c:\Users\16137\Desktop\vb期末\diagrams'
DOC_PATH = r'c:\Users\16137\Desktop\vb期末\VibeCoding期末报告_230152061_庞凯匀.docx'
OUT_PATH = r'c:\Users\16137\Desktop\vb期末\VibeCoding期末报告_230152061_庞凯匀_含图.docx'

# 映射：占位关键词 → PNG 文件路径
MAP = {
    '图1：系统三层架构图': 'fig1_architecture.png',
    '图2：核心数据流图': 'fig2_dataflow.png',
    '图3：系统功能模块图': 'fig5_modules.png',
    '图4：比喻匹配加权评分算法流程图': 'fig3_metaphor_match.png',
    # 图5-10 是界面截图，用户提供
}

doc = Document(DOC_PATH)

for p in doc.paragraphs:
    text = p.text.strip()
    for keyword, filename in MAP.items():
        if keyword in text:
            # Clear the paragraph text
            p.clear()
            p.paragraph_format.first_line_indent = Cm(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            img_path = os.path.join(D, filename)
            if os.path.exists(img_path):
                run = p.add_run()
                run.add_picture(img_path, width=Inches(5.5))
                # Add caption below
                caption_run = p.add_run(f'\n{keyword}')
                caption_run.font.size = Pt(9)
                caption_run.font.color.rgb = RGBColor(100, 116, 139)
                print(f'[OK] Embedded {filename} for {keyword}')
            else:
                print(f'[MISSING] {img_path}')
            break

doc.save(OUT_PATH)
print(f'\n[DONE] Updated: {OUT_PATH}')
print('Remaining: fig5-fig10 are UI screenshots — user needs to paste manually.')
